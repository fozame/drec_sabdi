"""
word.py
───────
Annexe Word (.docx, A4 paysage), directement insérable dans un rapport :
même structure que le tableau de référence (N° / Statistiques /
Spécifications / Résultats / Observations), cellules fusionnées, ligne
d'en-tête répétée sur chaque page, pagination en pied de page.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config
from . import modele as M

POLICE = "Arial"
ENCRE = RGBColor(0x1B, 0x24, 0x33)
GRIS = RGBColor(0x5B, 0x64, 0x72)
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
FOND_ENTETE = "DCE3EC"
FOND_SOUS_TOTAL = "EEF1F5"
LARGEUR_UTILE = 29.7 - 2 * 1.4


# ═════════════════════════════════════════════════════════════════════════════
# OUTILS XML
# ═════════════════════════════════════════════════════════════════════════════

def _ombre(cellule, couleur):
    tcPr = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tcPr.append(shd)


def _repeter_entete(ligne):
    trPr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _pas_de_coupure(ligne):
    trPr = ligne._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _bordures(table, epaisseur=4, couleur="9AA3AF"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(8 if cote in ("top", "left", "bottom", "right") else epaisseur))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "1B2433" if cote in ("top", "left", "bottom", "right") else couleur)
        b.append(e)
    tblPr.append(b)


def _marges_cellules(table, haut=30, gauche=60):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for cote, v in (("top", haut), ("bottom", haut), ("left", gauche), ("right", gauche)):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _largeur_fixe(table, largeurs_cm):
    """Largeurs de colonnes figées (grille + cellules), respectées par Word et LibreOffice."""
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(largeurs_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), largeurs_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for ligne in table.rows:
        for j, c in enumerate(ligne.cells):
            c.width = Cm(largeurs_cm[j])


def _champ(run, code):
    for typ, texte in (("begin", None), (None, code), ("end", None)):
        if typ:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), typ)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = texte
        run._r.append(e)


def _ecrire(cellule, texte, taille=7.5, gras=False, italique=False, couleur=ENCRE,
            align=WD_ALIGN_PARAGRAPH.LEFT, fond=None):
    cellule.text = ""
    lignes = str(texte if texte is not None else "").split("\n")
    p = cellule.paragraphs[0]
    for i, l in enumerate(lignes):
        if i:
            p = cellule.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        r = p.add_run(l)
        r.font.size = Pt(taille)
        r.font.bold = gras
        r.font.italic = italique
        r.font.color.rgb = couleur
        r.font.name = POLICE
    cellule.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fond:
        _ombre(cellule, fond)


def _para(doc, texte, taille=9, gras=False, couleur=ENCRE, align=WD_ALIGN_PARAGRAPH.LEFT,
          avant=0, apres=4, italique=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(avant)
    p.paragraph_format.space_after = Pt(apres)
    r = p.add_run(texte)
    r.font.size, r.font.bold, r.font.italic = Pt(taille), gras, italique
    r.font.color.rgb = couleur
    r.font.name = POLICE
    return p


# ═════════════════════════════════════════════════════════════════════════════
# DOCUMENT
# ═════════════════════════════════════════════════════════════════════════════

def _nouveau_document(entete_droite: str, paysage: bool = True) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = POLICE
    st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
    sec = doc.sections[0]
    if paysage:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.4) if paysage else Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.5)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.7)

    # En-tête
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Inches
    for nom_style in ("Header", "Footer"):
        doc.styles[nom_style].font.size = Pt(7)
        doc.styles[nom_style].font.name = POLICE
    hp = sec.header.paragraphs[0]
    hp.text = ""
    ts = hp.paragraph_format.tab_stops
    ts.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin, WD_TAB_ALIGNMENT.RIGHT)
    logo = config.chemin_logo()
    if logo:
        try:
            hp.add_run().add_picture(logo, height=Cm(0.8))
            hp.add_run("  ")
        except Exception:
            pass
    r = hp.add_run(config.ORGANISME[0])
    r.font.size, r.font.bold, r.font.color.rgb = Pt(7.5), True, GRIS
    r = hp.add_run("\t" + entete_droite)
    r.font.size, r.font.color.rgb = Pt(7.5), GRIS
    # Pied de page : Page x / y
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
    fp.paragraph_format.tab_stops.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin,
                                               WD_TAB_ALIGNMENT.RIGHT)
    r0 = fp.add_run(f"{config.NOM_APPLICATION} – {config.NOM_COMPLET}\t")
    r0.font.size, r0.font.color.rgb = Pt(7), GRIS
    for texte, code in (("Page ", None), (None, "PAGE"), (" / ", None), (None, "NUMPAGES")):
        run = fp.add_run(texte or "")
        run.font.size, run.font.color.rgb = Pt(7), GRIS
        if code:
            _champ(run, code)
    return doc


def _logo(doc, hauteur_cm=2.0):
    logo = config.chemin_logo()
    if not logo:
        return
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(logo, height=Cm(hauteur_cm))
    except Exception:
        pass


def _bloc_titre(doc, titre, sous_titre=None, mot="ANNEXE"):
    _logo(doc)
    _para(doc, config.ORGANISME[0], 9.5, True, align=WD_ALIGN_PARAGRAPH.CENTER, apres=0)
    for l in config.ORGANISME[1:]:
        _para(doc, l, 8.5, couleur=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, apres=0)
    _para(doc, mot, 13, True, ACCENT, WD_ALIGN_PARAGRAPH.CENTER, avant=10, apres=2)
    _para(doc, titre, 12, True, align=WD_ALIGN_PARAGRAPH.CENTER, apres=2)
    if sous_titre:
        _para(doc, sous_titre, 9, couleur=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, apres=10)


def _section(doc, texte):
    _para(doc, texte, 10.5, True, ACCENT, avant=10, apres=6)


def tableau_etat_docx(doc, res, numero_depart=1):
    mod = M.tableau_etat(res, numero_depart)
    cols = mod["colonnes"]
    k = len(cols)
    w_n, w_stat, w_spec, w_tot = 0.9, 5.0, 3.3, 2.0
    w_s = 2.0 if k <= 4 else 1.75
    w_obs = LARGEUR_UTILE - (w_n + w_stat + w_spec + w_tot + k * w_s)
    if w_obs < 4.5:
        w_stat -= 4.5 - w_obs
        w_obs = 4.5
    elif w_obs > 9.5:
        surplus = w_obs - 9.5
        w_stat += surplus * 0.45
        w_spec += surplus * 0.25
        w_s += surplus * 0.30 / k
        w_obs = 9.5
    largeurs = [w_n, w_stat, w_spec] + [w_s] * k + [w_tot, w_obs]
    nb_col = len(largeurs)

    t = doc.add_table(rows=2 + len(mod["lignes"]), cols=nb_col)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _largeur_fixe(t, largeurs)
    _bordures(t)
    _marges_cellules(t)

    # En-têtes (2 lignes)
    h1 = t.rows[1].cells
    C = WD_ALIGN_PARAGRAPH.CENTER
    for j, texte in ((0, "N°"), (1, "Statistiques"), (2, "Spécifications"), (nb_col - 1, "Observations")):
        m = t.cell(0, j).merge(t.cell(1, j))
        _ecrire(m, texte, 7.5, True, align=C, fond=FOND_ENTETE)
    m = t.cell(0, 3).merge(t.cell(0, 3 + k))
    _ecrire(m, "Résultats", 7.5, True, align=C, fond=FOND_ENTETE)
    for j, (_, lib) in enumerate(cols):
        _ecrire(h1[3 + j], lib, 7, True, align=C, fond=FOND_ENTETE)
    _ecrire(h1[3 + k], "Total", 7, True, align=C, fond=FOND_ENTETE)
    _repeter_entete(t.rows[0])
    _repeter_entete(t.rows[1])

    R_ = WD_ALIGN_PARAGRAPH.RIGHT
    r = 2
    for l in mod["lignes"]:
        st = l["style"]
        cells = t.rows[r].cells
        _pas_de_coupure(t.rows[r])
        fond = FOND_SOUS_TOTAL if st == "sous_total" else None
        gras = st == "sous_total"
        if l["debut_groupe"]:
            _ecrire(cells[0], l["n"], align=C, fond=fond)
            _ecrire(cells[1], l["statistique"], gras=gras, italique=st == "non_evaluable",
                    couleur=GRIS if st == "non_evaluable" else ENCRE, fond=fond)
            _ecrire(cells[nb_col - 1], l["observation"], 6.5, fond=fond)
        spec = l["specification"]
        _ecrire(cells[2], spec, gras=gras, fond=fond)
        if l["variantes"]:
            p = cells[2].add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(" | ".join(l["variantes"]))
            run.font.size, run.font.color.rgb, run.font.name = Pt(6), GRIS, POLICE
        for j, v in enumerate(l["valeurs"]):
            _ecrire(cells[3 + j], v, gras=gras, align=R_, fond=fond)
        _ecrire(cells[3 + k], l["total"], gras=True, align=R_, fond=fond)
        if l["debut_groupe"] and l["taille_groupe"] > 1:
            fin = r + l["taille_groupe"] - 1
            for j in (0, 1, nb_col - 1):
                texte_cellule = t.cell(r, j)
                texte_cellule.merge(t.cell(fin, j))
        r += 1
    return mod["numero_suivant"]


def _tableau_simple(doc, entetes, lignes, largeurs, droite=()):
    t = doc.add_table(rows=1 + len(lignes), cols=len(entetes))
    t.autofit = False
    _largeur_fixe(t, largeurs)
    _bordures(t)
    _marges_cellules(t)
    for j, e in enumerate(entetes):
        _ecrire(t.rows[0].cells[j], e, 7.5, True, align=WD_ALIGN_PARAGRAPH.CENTER, fond=FOND_ENTETE)
    _repeter_entete(t.rows[0])
    for i, l in enumerate(lignes, 1):
        _pas_de_coupure(t.rows[i])
        for j, v in enumerate(l):
            _ecrire(t.rows[i].cells[j], v, 7.5,
                    align=WD_ALIGN_PARAGRAPH.RIGHT if j in droite else WD_ALIGN_PARAGRAPH.LEFT)
    return t


def _controles(doc, res):
    lignes, famille = [], None
    for c in M.controles(res):
        lignes.append([c["famille"] if c["famille"] != famille else "", c["libelle"], c["valeur"],
                       c["base"], c["taux"], c["detail"] or "–"])
        famille = c["famille"]
    _tableau_simple(doc, ["Famille", "Contrôle", "Nombre", "Base", "Taux", "Détail"], lignes,
                    [3.0, 9.2, 2.0, 2.0, 1.6, LARGEUR_UTILE - 17.8], droite=(2, 3, 4))


def _signaux(doc, res):
    lignes, famille = [], None
    for x in M.signaux(res):
        det = "; ".join(f"{a} : {b}" for a, b in x["detail_txt"] if b is not None and b != "0")
        lignes.append([x["famille"] if x["famille"] != famille else "", x["libelle"], x["lignes_txt"],
                       x["entites_txt"] or "–", x["taux_txt"], det or "–"])
        famille = x["famille"]
    if lignes:
        _tableau_simple(doc, ["Famille", "Signal", "Numéros", "Entités", "Taux", "Détail"], lignes,
                        [2.6, 8.2, 2.0, 3.2, 1.6, LARGEUR_UTILE - 17.6], droite=(2, 3, 4))
    return bool(lignes)


def _trace(doc, res):
    for t in M.trace(res):
        _para(doc, t["titre"], 9, True, avant=6, apres=3)
        _tableau_simple(doc, ["Étape", "Lignes"], [[a, b] for a, b in t["etapes"]] + [[t["final_lib"], t["final"]]],
                        [13, 3], droite=(1,))


def _notes(doc, res):
    _para(doc, "Notes de méthode", 9, True, avant=8, apres=2)
    for n in M.notes_methode(res):
        _para(doc, "• " + n, 7.5, couleur=GRIS, apres=1)


def generer_docx(res: dict, contenu: str = "complet") -> bytes:
    doc = _nouveau_document(f"Annexe – {res.get('operateur_libelle', '')} – {res['periode']['libelle']}")
    _bloc_titre(doc, config.TITRE_ANNEXE, f"{res.get('operateur_libelle', '')} – mois de {res['periode']['libelle']}")
    _section(doc, M.titre_section(res, 0))
    tableau_etat_docx(doc, res)
    if contenu == "complet":
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if res.get("signaux"):
            _section(doc, "II. Signaux d'alerte")
            _para(doc, "Informations non visibles dans les chiffres déclarés. Âges et délais calculés au mois de la "
                       "base ; une même personne est reconnue par son nom et sa date de naissance.", 8, couleur=GRIS)
            _signaux(doc, res)
        _section(doc, "III. Contrôles complémentaires")
        _controles(doc, res)
        _section(doc, "IV. Complétude des champs d'identification")
        for bloc in M.completude(res):
            _para(doc, f"{bloc['libelle']} – effectif : {bloc['effectif']}", 9, True, avant=6, apres=3)
            _tableau_simple(doc, ["Champ", "Colonne du fichier", "Non renseignés", "Invalides", "Taux"],
                            [[l["libelle"], l["colonne"], l["vides"], l["invalides"], l["taux"]] for l in bloc["lignes"]],
                            [6, 6, 3, 3, 2.4], droite=(2, 3, 4))
        _section(doc, "V. Fichiers analysés et passage du fichier au tableau")
        _tableau_simple(doc, ["Fichier", "Rôle retenu", "Lignes", "Colonnes reconnues", "Statuts", "Champs non trouvés"],
                        [[f["nom"], f["role"], f["lignes"], f["reconnues"], f["statut"], f["manquants"]]
                         for f in M.fichiers(res)],
                        [6.8, 5.2, 2.2, 2.5, 2.0, LARGEUR_UTILE - 18.7], droite=(2,))
        _trace(doc, res)
    _notes(doc, res)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generer_docx_groupe(resultats: list[dict], contenu: str = "tableau", numerotation_continue=True) -> bytes:
    doc = _nouveau_document(f"Annexe – {config.TITRE_ANNEXE}")
    _bloc_titre(doc, config.TITRE_ANNEXE, " – ".join(
        f"{r.get('operateur_libelle', '')} ({r['periode']['libelle']})" for r in resultats))
    num = 1
    for i, res in enumerate(resultats):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _section(doc, M.titre_section(res, i))
        suivant = tableau_etat_docx(doc, res, num if numerotation_continue else 1)
        num = suivant
        if contenu == "complet":
            if res.get("signaux"):
                _para(doc, f"Signaux d'alerte – {res.get('operateur_libelle', '')}", 9, True, avant=8, apres=3)
                _signaux(doc, res)
            _para(doc, f"Contrôles complémentaires – {res.get('operateur_libelle', '')}", 9, True, avant=8, apres=3)
            _controles(doc, res)
    _notes(doc, resultats[0])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
